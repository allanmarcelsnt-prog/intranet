#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_json_treinamentos.py — Safenet Tecnologia
Le a planilha mais recente de treinamentos e publica data/treinamentos.json
para consumo via /api/dados/treinamentos.

Mesmo padrao do gerar_json.py (Comercial).
Rodar via cron: docker exec intranet-safenet python /app/gerar_json_treinamentos.py
"""

import json
import glob
import os
import sys
import unicodedata
from datetime import date, datetime

import openpyxl

# ---------------------------------------------------------------- config
PASTA_XLSX = "/data/treinamentos"          # volume montado do dataset Treinamentos
PADRAO     = "treinamentos_*.xlsx"
SAIDA      = "/app/data/treinamentos.json"

COLUNAS = [
    "fabricante", "curso", "data_inicio", "data_fim", "dias_texto",
    "hora_inicio", "hora_fim", "periodo", "publico", "restricao",
    "link_inscricao", "link_zoom", "obs_acesso", "gratuito", "modalidade",
]


def log(msg):
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def norm(txt):
    """minusculo, sem acento, sem espaco nas pontas."""
    if txt is None:
        return ""
    txt = str(txt).strip().lower()
    txt = unicodedata.normalize("NFKD", txt)
    return "".join(c for c in txt if not unicodedata.combining(c))


def parse_data(valor):
    """Aceita datetime/date do Excel ou string AAAA-MM-DD. Retorna date ou None."""
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def achar_planilha():
    arquivos = glob.glob(os.path.join(PASTA_XLSX, PADRAO))
    arquivos = [a for a in arquivos if not os.path.basename(a).startswith("~$")]
    if not arquivos:
        return None
    return max(arquivos, key=os.path.getmtime)


def main():
    caminho = achar_planilha()
    if not caminho:
        log(f"ERRO: nenhuma planilha '{PADRAO}' encontrada em {PASTA_XLSX}")
        sys.exit(1)

    log(f"Lendo: {os.path.basename(caminho)}")
    wb = openpyxl.load_workbook(caminho, data_only=True)

    # aba de dados: primeira que nao seja LEIA-ME
    ws = None
    for aba in wb.worksheets:
        if norm(aba.title) not in ("leia-me", "leiame", "leia me"):
            ws = aba
            break
    if ws is None:
        log("ERRO: nenhuma aba de dados encontrada")
        sys.exit(1)

    # mapa cabecalho -> indice, tolerante a acento/caixa/ordem
    cabecalho = [norm(c.value) for c in ws[1]]
    idx = {}
    for col in COLUNAS:
        if col in cabecalho:
            idx[col] = cabecalho.index(col)

    faltando = [c for c in ("curso", "data_inicio") if c not in idx]
    if faltando:
        log(f"ERRO: colunas obrigatorias ausentes: {faltando}")
        sys.exit(1)

    hoje = date.today()
    cursos = []
    ignorados = 0

    for linha in ws.iter_rows(min_row=2, values_only=True):
        def val(col):
            i = idx.get(col)
            if i is None or i >= len(linha):
                return ""
            v = linha[i]
            return "" if v is None else str(v).strip()

        curso = val("curso")
        if not curso:
            continue

        d_ini = parse_data(linha[idx["data_inicio"]])
        if d_ini is None:
            log(f"  IGNORADO (data invalida): {curso}")
            ignorados += 1
            continue

        d_fim = parse_data(linha[idx["data_fim"]]) if "data_fim" in idx else None
        d_fim = d_fim or d_ini

        # REGRA: some sozinho depois do ultimo dia do curso
        if d_fim < hoje:
            ignorados += 1
            continue

        registro = {
            "fabricante":     val("fabricante") or "Dahua",
            "curso":          curso,
            "data_inicio":    d_ini.isoformat(),
            "data_fim":       d_fim.isoformat(),
            "dias_texto":     val("dias_texto"),
            "hora_inicio":    val("hora_inicio")[:5],
            "hora_fim":       val("hora_fim")[:5],
            "periodo":        val("periodo"),
            "publico":        val("publico") or "Ambos",
            "restricao":      val("restricao"),
            "link_inscricao": val("link_inscricao"),
            "link_zoom":      val("link_zoom"),
            "obs_acesso":     val("obs_acesso"),
            "gratuito":       norm(val("gratuito")) in ("sim", "s", "true", "1", "x"),
            "modalidade":     val("modalidade"),
            "dias_restantes": (d_ini - hoje).days,
            "em_andamento":   d_ini <= hoje <= d_fim,
        }
        cursos.append(registro)

    cursos.sort(key=lambda c: c["data_inicio"])

    saida = {
        "atualizado_em": datetime.now().isoformat(timespec="seconds"),
        "origem":        os.path.basename(caminho),
        "total":         len(cursos),
        "proximo":       cursos[0] if cursos else None,
        "cursos":        cursos,
    }

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    tmp = SAIDA + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)
    os.replace(tmp, SAIDA)   # escrita atomica: o portal nunca le arquivo pela metade

    log(f"OK -> {SAIDA}")
    log(f"   {len(cursos)} curso(s) ativo(s), {ignorados} ignorado(s)/encerrado(s)")
    if cursos:
        p = cursos[0]
        log(f"   proximo: {p['curso']} ({p['data_inicio']}, faltam {p['dias_restantes']} dia(s))")


if __name__ == "__main__":
    main()
